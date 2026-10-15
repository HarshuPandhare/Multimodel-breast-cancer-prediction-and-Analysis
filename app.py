import os
import re
import time
import json
import uuid
import joblib
import pandas as pd
import torch
import torch.nn as nn
from flask import Flask, render_template, request, send_from_directory, jsonify, send_file
from torchvision import models, transforms
from PIL import Image
from werkzeug.utils import secure_filename
from dotenv import load_dotenv
from google import genai
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

load_dotenv()

app = Flask(__name__)

CLINICAL_MODEL_PATH = "models/breast_cancer_model.pkl"
IMAGE_MODEL_PATH = "models/breast_cancer_efficientnet.pth"
UPLOAD_FOLDER = "uploads"
REPORT_FOLDER = "reports"
ALLOWED_EXTENSIONS = {"png", "jpg", "jpeg"}

app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(REPORT_FOLDER, exist_ok=True)

device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
print("Using device:", device)
if torch.cuda.is_available():
    print("GPU:", torch.cuda.get_device_name(0))

if not os.path.exists(CLINICAL_MODEL_PATH):
    raise FileNotFoundError("Clinical model not found: " + CLINICAL_MODEL_PATH)

clinical_package = joblib.load(CLINICAL_MODEL_PATH)
clinical_model = clinical_package["model"]
target_encoder = clinical_package["target_encoder"]
clinical_accuracy = clinical_package["accuracy"]

if not os.path.exists(IMAGE_MODEL_PATH):
    raise FileNotFoundError("Image model not found: " + IMAGE_MODEL_PATH)

image_model = models.efficientnet_b0(weights=None)
number_of_features = image_model.classifier[1].in_features
image_model.classifier[1] = nn.Linear(number_of_features, 2)

checkpoint = torch.load(IMAGE_MODEL_PATH, map_location=device)
image_model.load_state_dict(checkpoint["model_state_dict"])
image_model = image_model.to(device)
image_model.eval()

image_classes = checkpoint["classes"] if "classes" in checkpoint else ["benign", "malignant"]
print("Image classes:", image_classes)

image_transform = transforms.Compose([
    transforms.Resize((224, 224)),
    transforms.ToTensor(),
    transforms.Normalize(mean=[0.485, 0.456, 0.406], std=[0.229, 0.224, 0.225])
])

# --- Gemini AI Agent Setup ---
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.8-flash")
GEMINI_FALLBACK_MODEL = os.getenv("GEMINI_FALLBACK_MODEL", "gemini-3.7-flash")

gemini_client = None
if GEMINI_API_KEY:
    try:
        gemini_client = genai.Client(api_key=GEMINI_API_KEY)
        print(f"Gemini AI agent initialized with model: {GEMINI_MODEL}")
    except Exception as e:
        print(f"Warning: Failed to initialize Gemini client: {e}")


def encode_marker(value):
    value = str(value).strip().lower()
    if value in ["positive", "pos", "1"]:
        return 1
    if value in ["negative", "neg", "0"]:
        return 0
    raise ValueError("Invalid biomarker value.")


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def predict_image(image_path):
    image = Image.open(image_path).convert("RGB")
    image_tensor = image_transform(image).unsqueeze(0).to(device)

    with torch.no_grad():
        outputs = image_model(image_tensor)
        probabilities = torch.softmax(outputs, dim=1)
        confidence, predicted = torch.max(probabilities, 1)

    predicted_class = image_classes[predicted.item()]
    confidence_value = confidence.item() * 100

    class_probabilities = []
    for i, class_name in enumerate(image_classes):
        class_probabilities.append({
            "class": class_name,
            "probability": round(probabilities[0][i].item() * 100, 2)
        })

    return predicted_class, round(confidence_value, 2), class_probabilities


def build_gemini_prompt(patient_data, clinical_result, image_result):
    """Build a structured prompt for the Gemini AI agent."""
    prompt = """You are a medical AI documentation assistant specializing in breast cancer analysis.
Based on the following patient data and ML model predictions, generate a comprehensive clinical documentation report.

IMPORTANT: You must respond with ONLY valid JSON, no markdown formatting, no code blocks, no extra text.
Use the exact JSON structure shown below.

PATIENT DATA:
- Age at Diagnosis: {age}
- ER Status: {er}
- PR Status: {pr}
- HER2 Status: {her2}
- Tumor Size: {tumor_size} mm
- Positive Lymph Nodes: {lymph_nodes}

ML MODEL PREDICTIONS:
- Clinical/Biomarker Model Prediction (PAM50 Subtype): {clinical_prediction}
- Clinical Model Confidence: {clinical_confidence}%
""".format(
        age=patient_data.get("age", "N/A"),
        er=patient_data.get("er_status", "N/A"),
        pr=patient_data.get("pr_status", "N/A"),
        her2=patient_data.get("her2_status", "N/A"),
        tumor_size=patient_data.get("tumor_size", "N/A"),
        lymph_nodes=patient_data.get("lymph_nodes", "N/A"),
        clinical_prediction=clinical_result.get("prediction", "N/A"),
        clinical_confidence=clinical_result.get("confidence", "N/A")
    )

    if image_result:
        prompt += """- Ultrasound Image Model Prediction: {image_prediction}
- Image Model Confidence: {image_confidence}%
""".format(
            image_prediction=image_result.get("prediction", "N/A"),
            image_confidence=image_result.get("confidence", "N/A")
        )

    prompt += """
Respond ONLY with this JSON structure (no markdown, no code fences):
{
    "summary": "A 2-3 sentence executive summary of the patient case and key findings.",
    "risk_assessment": "Detailed risk assessment based on the biomarkers and model predictions. Include risk level (Low/Moderate/High/Very High) and reasoning.",
    "subtype_analysis": "Explanation of the predicted PAM50 subtype, its characteristics, typical prognosis, and treatment response patterns.",
    "biomarker_interpretation": "Interpretation of ER, PR, and HER2 status combination and what it means for treatment selection.",
    "recommended_actions": [
        "First recommended action with explanation",
        "Second recommended action with explanation",
        "Third recommended action with explanation",
        "Fourth recommended action with explanation",
        "Fifth recommended action with explanation"
    ],
    "additional_tests": [
        "First additional test recommendation",
        "Second additional test recommendation",
        "Third additional test recommendation"
    ],
    "treatment_considerations": "Overview of potential treatment approaches based on the molecular subtype and biomarker profile.",
    "follow_up_plan": "Suggested follow-up schedule and monitoring plan.",
    "confidence_notes": "Notes about the AI model confidence levels and any caveats about the predictions."
}

IMPORTANT GUIDELINES:
- This is for research/educational purposes only - always include appropriate disclaimers
- Be thorough and evidence-based in your analysis
- Consider the interplay between all biomarkers and clinical features
- Provide actionable, specific recommendations
- Do NOT wrap the response in markdown code blocks or any other formatting
"""
    return prompt


def call_gemini_agent(prompt):
    """Call the Gemini AI agent with fallback model support."""
    if not gemini_client:
        raise ValueError("Gemini AI agent not configured. Check your GEMINI_API_KEY in .env")

    models_to_try = [GEMINI_MODEL, GEMINI_FALLBACK_MODEL]

    for model_name in models_to_try:
        try:
            response = gemini_client.models.generate_content(
                model=model_name,
                contents=prompt
            )
            return response.text
        except Exception as e:
            print(f"Gemini model {model_name} failed: {e}")
            continue

    raise RuntimeError("All Gemini models failed. Please check your API key and model names.")


def parse_gemini_response(response_text):
    """Parse the JSON response from Gemini, handling markdown code fences."""
    cleaned = response_text.strip()

    # Remove markdown code fences if present
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?\s*\n?", "", cleaned)
        cleaned = re.sub(r"\n?```\s*$", "", cleaned)
        cleaned = cleaned.strip()

    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        # Try to extract JSON from the response
        json_match = re.search(r'\{[\s\S]*\}', cleaned)
        if json_match:
            try:
                return json.loads(json_match.group())
            except json.JSONDecodeError:
                pass

        return {
            "summary": response_text[:500],
            "risk_assessment": "Could not parse structured response from AI agent.",
            "subtype_analysis": "",
            "biomarker_interpretation": "",
            "recommended_actions": ["Please regenerate the report for structured recommendations."],
            "additional_tests": [],
            "treatment_considerations": "",
            "follow_up_plan": "",
            "confidence_notes": "Response parsing failed - raw AI response was preserved in summary."
        }


def generate_docx_report(patient_data, clinical_result, image_result, ai_analysis, report_id):
    """Generate a professional .docx report with all analysis data."""
    doc = Document()

    # Title
    title = doc.add_heading("Breast Cancer AI Analysis Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = RGBColor(37, 99, 235)

    doc.add_paragraph(f"Report ID: {report_id}")
    doc.add_paragraph(f"Generated: {time.strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph("")

    # Patient Information
    doc.add_heading("Patient Information", level=1)
    table = doc.add_table(rows=6, cols=2)
    table.style = "Light Grid Accent 1"
    fields = [
        ("Age at Diagnosis", f"{patient_data.get('age', 'N/A')}"),
        ("ER Status", f"{patient_data.get('er_status', 'N/A')}"),
        ("PR Status", f"{patient_data.get('pr_status', 'N/A')}"),
        ("HER2 Status", f"{patient_data.get('her2_status', 'N/A')}"),
        ("Tumor Size", f"{patient_data.get('tumor_size', 'N/A')} mm"),
        ("Positive Lymph Nodes", f"{patient_data.get('lymph_nodes', 'N/A')}")
    ]
    for i, (label, value) in enumerate(fields):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = str(value)

    doc.add_paragraph("")

    # ML Predictions
    doc.add_heading("ML Model Predictions", level=1)

    doc.add_heading("Clinical/Biomarker Model", level=2)
    doc.add_paragraph(f"Predicted PAM50 Subtype: {clinical_result.get('prediction', 'N/A')}")
    doc.add_paragraph(f"Model Confidence: {clinical_result.get('confidence', 'N/A')}%")

    if image_result:
        doc.add_heading("Ultrasound Image Model", level=2)
        doc.add_paragraph(f"Image Prediction: {image_result.get('prediction', 'N/A')}")
        doc.add_paragraph(f"Model Confidence: {image_result.get('confidence', 'N/A')}%")

    doc.add_paragraph("")

    # AI Agent Analysis
    doc.add_heading("AI Agent Analysis & Documentation", level=1)

    if ai_analysis.get("summary"):
        doc.add_heading("Executive Summary", level=2)
        doc.add_paragraph(ai_analysis["summary"])

    if ai_analysis.get("risk_assessment"):
        doc.add_heading("Risk Assessment", level=2)
        doc.add_paragraph(ai_analysis["risk_assessment"])

    if ai_analysis.get("subtype_analysis"):
        doc.add_heading("Subtype Analysis", level=2)
        doc.add_paragraph(ai_analysis["subtype_analysis"])

    if ai_analysis.get("biomarker_interpretation"):
        doc.add_heading("Biomarker Interpretation", level=2)
        doc.add_paragraph(ai_analysis["biomarker_interpretation"])

    if ai_analysis.get("treatment_considerations"):
        doc.add_heading("Treatment Considerations", level=2)
        doc.add_paragraph(ai_analysis["treatment_considerations"])

    if ai_analysis.get("recommended_actions"):
        doc.add_heading("Recommended Actions", level=2)
        for action in ai_analysis["recommended_actions"]:
            doc.add_paragraph(action, style="List Bullet")

    if ai_analysis.get("additional_tests"):
        doc.add_heading("Additional Tests Recommended", level=2)
        for test in ai_analysis["additional_tests"]:
            doc.add_paragraph(test, style="List Bullet")

    if ai_analysis.get("follow_up_plan"):
        doc.add_heading("Follow-Up Plan", level=2)
        doc.add_paragraph(ai_analysis["follow_up_plan"])

    if ai_analysis.get("confidence_notes"):
        doc.add_heading("Confidence & Limitations", level=2)
        doc.add_paragraph(ai_analysis["confidence_notes"])

    # Disclaimer
    doc.add_paragraph("")
    disclaimer = doc.add_paragraph()
    disclaimer_run = disclaimer.add_run(
        "DISCLAIMER: This report is generated by an AI system for research and educational purposes only. "
        "It should not be used as a substitute for professional medical diagnosis or treatment. "
        "Always consult with qualified healthcare professionals for medical decisions."
    )
    disclaimer_run.font.size = Pt(9)
    disclaimer_run.font.italic = True
    disclaimer_run.font.color.rgb = RGBColor(128, 128, 128)

    # Save
    filepath = os.path.join(REPORT_FOLDER, f"breast_cancer_report_{report_id}.docx")
    doc.save(filepath)
    return filepath


@app.route("/", methods=["GET", "POST"])
def home():
    clinical_prediction = None
    clinical_confidence = None
    clinical_probabilities = None
    image_prediction = None
    image_confidence = None
    image_probabilities = None
    uploaded_image = None
    error = None
    report = None

    if request.method == "POST":
        try:
            er_status = request.form.get("er_status")
            pr_status = request.form.get("pr_status")
            her2_status = request.form.get("her2_status")
            age = float(request.form.get("age"))
            tumor_size = float(request.form.get("tumor_size"))
            lymph_nodes = float(request.form.get("lymph_nodes"))

            if age <= 0:
                raise ValueError("Age must be greater than 0.")
            if tumor_size < 0:
                raise ValueError("Tumor size cannot be negative.")
            if lymph_nodes < 0:
                raise ValueError("Lymph node count cannot be negative.")

            input_data = pd.DataFrame({
                "er_status": [encode_marker(er_status)],
                "pr_status": [encode_marker(pr_status)],
                "her2_status": [encode_marker(her2_status)],
                "age_at_diagnosis": [age],
                "tumor_size": [tumor_size],
                "lymph_nodes_examined_positive": [lymph_nodes]
            })

            prediction_encoded = clinical_model.predict(input_data)[0]
            clinical_prediction = target_encoder.inverse_transform([prediction_encoded])[0]

            probability_array = clinical_model.predict_proba(input_data)[0]
            clinical_confidence = round(float(max(probability_array)) * 100, 2)

            clinical_probabilities = []
            for class_name, probability in zip(target_encoder.classes_, probability_array):
                clinical_probabilities.append({
                    "class": class_name,
                    "probability": round(float(probability) * 100, 2)
                })
            clinical_probabilities.sort(key=lambda x: x["probability"], reverse=True)

            report = {
                "ER": er_status,
                "PR": pr_status,
                "HER2": her2_status,
                "Age": age,
                "Tumor Size": tumor_size,
                "Lymph Nodes": lymph_nodes
            }

            image_file = request.files.get("image")
            if image_file:
                if image_file.filename == "":
                    raise ValueError("Please select an image.")
                if not allowed_file(image_file.filename):
                    raise ValueError("Only PNG, JPG and JPEG images are allowed.")

                filename = str(int(time.time())) + "_" + secure_filename(image_file.filename)
                image_path = os.path.join(app.config["UPLOAD_FOLDER"], filename)
                image_file.save(image_path)
                uploaded_image = "/uploads/" + filename

                image_prediction, image_confidence, image_probabilities = predict_image(image_path)

        except Exception as e:
            error = str(e)

    return render_template(
        "index.html",
        clinical_prediction=clinical_prediction,
        clinical_confidence=clinical_confidence,
        clinical_probabilities=clinical_probabilities,
        image_prediction=image_prediction,
        image_confidence=image_confidence,
        image_probabilities=image_probabilities,
        uploaded_image=uploaded_image,
        error=error,
        report=report,
        clinical_accuracy=round(clinical_accuracy * 100, 2)
    )


@app.route("/generate-report", methods=["POST"])
def generate_report():
    """API endpoint to generate AI documentation report."""
    try:
        data = request.get_json()
        if not data:
            return jsonify({"error": "No data provided"}), 400

        patient_data = data.get("patient_data", {})
        clinical_result = data.get("clinical_result", {})
        image_result = data.get("image_result", None)

        # Build prompt and call Gemini
        prompt = build_gemini_prompt(patient_data, clinical_result, image_result)
        raw_response = call_gemini_agent(prompt)
        ai_analysis = parse_gemini_response(raw_response)

        # Generate .docx report
        report_id = uuid.uuid4().hex
        report_path = generate_docx_report(
            patient_data, clinical_result, image_result, ai_analysis, report_id
        )

        return jsonify({
            "success": True,
            "report_id": report_id,
            "ai_analysis": ai_analysis,
            "download_url": f"/download-report/{report_id}"
        })

    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/download-report/<report_id>")
def download_report(report_id):
    """Download a generated .docx report."""
    # Sanitize report_id to prevent path traversal
    safe_id = secure_filename(report_id)
    filename = f"breast_cancer_report_{safe_id}.docx"
    filepath = os.path.join(REPORT_FOLDER, filename)

    if not os.path.exists(filepath):
        return "Report not found", 404

    return send_file(filepath, as_attachment=True, download_name=filename)


@app.route("/uploads/<filename>")
def uploaded_file(filename):
    return send_from_directory(app.config["UPLOAD_FOLDER"], filename)


if __name__ == "__main__":
    app.run(debug=True, host="127.0.0.1", port=5000)